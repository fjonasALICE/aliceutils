#!/usr/bin/env python3
"""
alicemeetingassistant.py - Answer questions about ALICE meetings on Indico with an AI agent.

Usage:
    python alicemeetingassistant.py "What's going on in the ALICE Physics Forum?"
    python alicemeetingassistant.py "Summarise the last PWG-LF meeting"
    python alicemeetingassistant.py "Contributions of Jane Doe in https://indico.cern.ch/event/1234567/"
    python alicemeetingassistant.py --chat
    python alicemeetingassistant.py "<question>" --config alicemeetingassistant.json
"""

import os
import sys
import subprocess

# ---------------------------------------------------------------------------
# Bootstrap: ensure we're running inside a venv with pydantic-ai + requests
# ---------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
VENV_DIR = os.path.join(SCRIPT_DIR, ".venv_alicemeetingassistant")
DEPENDENCIES = ["pydantic-ai", "requests", "pypdf", "python-pptx"]


def _bootstrap() -> None:
    """If not already inside the project venv, create it and re-exec."""
    venv_python = os.path.join(VENV_DIR, "bin", "python")
    if os.path.abspath(sys.executable) == os.path.abspath(venv_python):
        return

    if not os.path.isfile(venv_python):
        import venv as _venv

        print(f"[setup] Creating virtual environment at {VENV_DIR} …")
        _venv.create(VENV_DIR, with_pip=True)
        pip = os.path.join(VENV_DIR, "bin", "pip")
        print(f"[setup] Installing dependencies ({', '.join(DEPENDENCIES)}) …")
        subprocess.check_call([pip, "install", "--quiet", *DEPENDENCIES])
        print("[setup] Done. Re-starting inside venv …\n")

    os.execv(venv_python, [venv_python] + sys.argv)


_bootstrap()

# ---------------------------------------------------------------------------
# Main imports (only reached when running inside the venv)
# ---------------------------------------------------------------------------
import argparse
import datetime
import difflib
import hashlib
import io
import json
import re
import tempfile
import unicodedata
from pathlib import Path

os.environ.setdefault("PYDANTIC_AI_NO_BANNER", "1")

import requests
from pydantic_ai import Agent, UsageLimits
from pydantic_ai.exceptions import ModelHTTPError, UsageLimitExceeded
from pydantic_ai.models.openai import OpenAIChatModel
from pydantic_ai.models.openrouter import OpenRouterModel
from pydantic_ai.profiles.openai import OpenAIModelProfile
from pydantic_ai.providers.openai import OpenAIProvider
from pydantic_ai.providers.openrouter import OpenRouterProvider

INDICO = "https://indico.cern.ch"
ALICE_MEETINGS = 6746
MAX_EVENTS = 150
MAX_HITS = 100
MAX_CHARS = 15_000
MAX_PAGES = 15
MAX_BYTES = 50_000_000
LINE_CLIP = 300
HEADERS = {"User-Agent": "aliceutils-alicemeetingassistant"}

MIN_PROXY = 0.5
MAX_PROXIES = 15
SEARCH_PAGES = 5
STOPWORDS = {"and", "of", "the", "for", "in", "on", "a", "meeting", "meetings", "alice"}

_DATE_RE = re.compile(r"today|now|[+-]?\d{1,4}d|\d{4}-\d{2}-\d{2}")

INSTRUCTIONS = (
    "You answer questions about ALICE collaboration meetings on CERN Indico (indico.cern.ch). "
    "Use the tools to read Indico. Do not invent meetings, talks, speakers or links. "
    f"ALICE Meetings is category {ALICE_MEETINGS}. Listing it also lists all its subcategories. "
    "Known subcategories: Physics (302), Physics Forum (6884). "
    "For other groups (a PWG, a PAG, a detector, a task force), call find_category, "
    "or list ALICE Meetings and filter by query. "
    "Group names on Indico often differ from how people say them: acronyms, abbreviations, "
    "reordered words, or renamed groups. For example 'Nuclei, Exotica and Spectra PAG' may be "
    "'NuSpEx PAG', and 'Light Flavour' may be 'PWG-LF'. If there is no exact match, use the "
    "closest proxy the tools return, try other spellings or acronyms yourself if needed, "
    "and tell the user which category or meeting you used in place of the name they gave. "
    "When asked what is going on in a forum or group, list the recent and upcoming meetings, "
    "open the relevant ones with get_meeting, and give a breakdown per meeting: date, "
    "then each contribution with its title, speakers, and links to the slides. "
    "Read slides with read_attachment when the question asks about content, results or conclusions, "
    "or when a title alone does not say what a talk is about. Read only the pages you need. "
    "If slides cannot be read, say so and quote the reason the tool gives, word for word. "
    "To find the contributions of one person, call find_contributions. "
    "Pass the event when the question names a specific meeting. "
    "Always give Indico links: the meeting URL, and the attachment URL for every slide you mention. "
    "Format the answer in Markdown. Keep it concise and organised by meeting."
)


class Context:
    """Indico session, caches for event data and downloaded attachments."""

    def __init__(self, token: str | None) -> None:
        self.session = requests.Session()
        self.session.headers.update(HEADERS)
        if token:
            self.session.headers["Authorization"] = f"Bearer {token}"
        self.token = token
        self.cache = Path(tempfile.mkdtemp(prefix="alicemeetingassistant_"))
        self.events: dict[str, dict] = {}
        self.docs: dict[str, list[str]] = {}
        self.active: dict[int, str] | None = None

    def get_json(self, path: str, params: dict | None = None) -> dict:
        try:
            resp = self.session.get(INDICO + path, params=params, timeout=120)
        except requests.RequestException as exc:
            raise RuntimeError(f"Indico request failed: {exc}") from exc
        if resp.status_code in (401, 403):
            raise RuntimeError(f"Indico refused {path} ({resp.status_code}). {self.auth_hint()}")
        if resp.status_code != 200:
            raise RuntimeError(f"Indico {path} returned {resp.status_code}")
        if "json" not in resp.headers.get("content-type", ""):
            raise RuntimeError(f"Indico {path} did not return JSON. {self.auth_hint()}")
        return resp.json()

    def export(self, what: str, params: dict) -> list[dict]:
        data = self.get_json(f"/export/{what}.json", params)
        return data.get("results") or []

    def auth_hint(self) -> str:
        if not self.token:
            return "No Indico token is set, so only public meetings are visible."
        return "The Indico token may be invalid, expired, or missing a scope."


QUIET = False


def log(message: str) -> None:
    """Progress line for one-shot runs; muted in the chat UI."""
    if not QUIET:
        print(f"  {message}", flush=True)


def clip(text: str, size: int = LINE_CLIP) -> str:
    text = " ".join(text.split())
    return text if len(text) <= size else text[:size] + " …"


def fold(text: str) -> str:
    """Lowercase and strip accents, for name and title matching."""
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch)).lower()


def words(text: str) -> list[str]:
    """Folded words, with CamelCase split too: 'NuSpEx PAG' -> nuspex, nu, sp, ex, pag."""
    out: list[str] = []
    for raw in re.findall(r"[A-Za-z0-9]+", unicodedata.normalize("NFKD", text)):
        word = fold(raw)
        if word in STOPWORDS:
            continue
        out.append(word)
        parts = re.findall(r"[A-Z][a-z]+|[a-z]+|[0-9]+", raw)
        if len(parts) > 1:
            out += [fold(p) for p in parts if fold(p) not in STOPWORDS]
    return out


def abbreviates(short: str, full: list[str]) -> set[str]:
    """Words of full that short abbreviates. Either the initials of consecutive words
    ('lf' covers light, flavour), or prefixes of 2+ letters in any order
    ('nuspex' covers nuclei, spectra, exotica)."""
    if not 2 <= len(short) <= 8:
        return set()
    for start in range(len(full) - len(short) + 1):
        window = full[start : start + len(short)]
        if "".join(w[0] for w in window) == short:
            return set(window)

    def split(rest: str, used: frozenset[str]) -> set[str] | None:
        if not rest:
            return set(used)
        for i in range(len(rest), 1, -1):
            for word in full:
                if word not in used and len(word) > i and word.startswith(rest[:i]):
                    found = split(rest[i:], used | {word})
                    if found:
                        return found
        return None

    found = split(short, frozenset())
    return found if found and len(found) >= 2 else set()


def proxy_score(query: str, text: str) -> float:
    """Fraction of the query words that text matches, allowing prefixes, typos and acronyms."""
    qwords = list(dict.fromkeys(words(query)))
    twords = list(dict.fromkeys(words(text)))
    if not qwords or not twords:
        return 0.0
    covered: set[str] = set()
    for q in qwords:
        for t in twords:
            if q == t or (min(len(q), len(t)) >= 3 and (q.startswith(t) or t.startswith(q))):
                covered.add(q)
            elif min(len(q), len(t)) >= 5 and difflib.SequenceMatcher(None, q, t).ratio() >= 0.85:
                covered.add(q)
    for t in twords:
        covered |= abbreviates(t, qwords)
    for q in qwords:
        if abbreviates(q, twords):
            covered.add(q)
    return len(covered) / len(qwords)


def check_date(value: str, name: str) -> str:
    value = value.strip()
    if not _DATE_RE.fullmatch(value):
        raise ValueError(f"{name} must be today, -30d, +7d or YYYY-MM-DD, got {value!r}")
    return value


def event_id(event: str) -> str:
    """Event id from an id or an Indico event URL."""
    event = str(event).strip()
    match = re.search(r"/event/(\d+)", event)
    if match:
        return match.group(1)
    if event.isdigit():
        return event
    raise ValueError(f"Not an Indico event id or URL: {event!r}")


def when(item: dict) -> str:
    start = item.get("startDate") or {}
    date, time = start.get("date", "?"), (start.get("time") or "")[:5]
    return f"{date} {time}".strip()


def person(p: dict) -> str:
    name = f"{p.get('first_name', '')} {p.get('last_name', '')}".strip() or p.get("fullName", "?")
    aff = p.get("affiliation")
    return f"{name} ({aff})" if aff else name


def attachments(item: dict) -> list[str]:
    lines: list[str] = []
    for folder in item.get("folders") or []:
        for att in folder.get("attachments") or []:
            url = att.get("download_url") or att.get("link_url") or ""
            size = att.get("size")
            extra = f", {size / 1e6:.1f} MB" if size else ""
            kind = "link" if att.get("type") == "link" else att.get("content_type") or "file"
            lines.append(f"    - {att.get('title') or att.get('filename')} [{kind}{extra}] {url}")
    return lines


def contribution_people(contrib: dict) -> list[tuple[str, dict]]:
    out = [("speaker", p) for p in contrib.get("speakers") or []]
    out += [("author", p) for p in contrib.get("primaryauthors") or []]
    out += [("co-author", p) for p in contrib.get("coauthors") or []]
    return out


def format_contribution(contrib: dict) -> str:
    speakers = ", ".join(person(p) for p in contrib.get("speakers") or []) or "no speaker listed"
    session = f" [session: {contrib['session']}]" if contrib.get("session") else ""
    duration = f" ({contrib['duration']} min)" if contrib.get("duration") else ""
    lines = [
        f"- {when(contrib)}{duration} {contrib.get('title', '?')}{session}",
        f"    speakers: {speakers}",
        f"    url: {contrib.get('url', '')}",
    ]
    if contrib.get("description"):
        lines.append(f"    description: {clip(contrib['description'])}")
    files = attachments(contrib)
    lines += files if files else ["    (no attachments)"]
    return "\n".join(lines)


def load_event(ctx: Context, event: str) -> dict:
    eid = event_id(event)
    if eid not in ctx.events:
        results = ctx.export(f"event/{eid}", {"detail": "contributions"})
        if not results:
            raise RuntimeError(f"Event {eid} not found or not visible. {ctx.auth_hint()}")
        ctx.events[eid] = results[0]
    return ctx.events[eid]


def active_categories(ctx: Context) -> dict[int, str]:
    """Category paths of every ALICE meeting in the last year, from one export call."""
    if ctx.active is None:
        data = ctx.get_json(
            f"/export/categ/{ALICE_MEETINGS}.json", {"from": "-365d", "to": "+60d"}
        )
        ctx.active = {}
        info = (data.get("additionalInfo") or {}).get("eventCategories") or []
        for item in info:
            names = [p.get("name", "") for p in item.get("path") or []]
            ctx.active[int(item["categoryId"])] = " / ".join(names)
        for ev in data.get("results") or []:
            ctx.active.setdefault(int(ev["categoryId"]), ev.get("category", ""))
    return ctx.active


def tool_find_category(ctx: Context, query: str) -> str:
    log(f"find_category {query!r}")
    query = query.strip()
    found: dict[int, str] = {}
    errors: list[str] = []
    # The full name first, then each word alone: Indico only matches substrings.
    terms = [query] + [w for w in dict.fromkeys(words(query)) if len(w) >= 3][:6]
    for term in terms:
        page, pages = 1, 1
        while page <= min(pages, SEARCH_PAGES):
            try:
                data = ctx.get_json("/category/search", {"q": term, "page": page})
            except RuntimeError as exc:
                errors.append(str(exc))
                break
            cats = data.get("categories") or []
            for cat in cats:
                path = " / ".join(p["title"] for p in cat.get("path") or [])
                if "ALICE" in path:
                    found[int(cat["id"])] = path
            if not cats:
                break
            pages = -(-int(data.get("total_count") or 0) // len(cats))
            page += 1
    try:
        for cid, path in active_categories(ctx).items():
            found.setdefault(cid, path)
    except RuntimeError as exc:
        errors.append(str(exc))
    if not found:
        return "\n".join(errors) or f"no ALICE categories match {query!r}"

    needle = fold(query)

    def rank(item: tuple[int, str]) -> tuple:
        path = item[1]
        title = path.rsplit(" / ", 1)[-1]
        exact = needle in fold(title)
        # The title counts most; the parent path breaks ties (e.g. PWG-LF for a PAG).
        score = max(proxy_score(query, title), 0.9 * proxy_score(query, path))
        return (not exact, -score, "ARCHIVE" in path.upper(), len(path))

    ranked = sorted(found.items(), key=rank)
    lines = []
    for cid, path in ranked[:MAX_PROXIES]:
        title = path.rsplit(" / ", 1)[-1]
        exact = needle in fold(title)
        score = max(proxy_score(query, title), 0.9 * proxy_score(query, path))
        if not exact and score < MIN_PROXY:
            continue
        tag = "exact" if exact else f"proxy {score:.2f}"
        lines.append(f"{cid}: {path} [{tag}]")
    if not lines:
        return f"no ALICE category resembles {query!r}"
    if "[exact]" not in lines[0]:
        lines.insert(0, f"No category is named {query!r}. Closest proxies "
                        "(acronyms, abbreviations, reordered words):")
    return "\n".join(lines)


def tool_list_meetings(ctx: Context, category_id: int, start: str, end: str, query: str) -> str:
    log(f"list_meetings {category_id} {start}..{end} {query!r}")
    try:
        params = {"from": check_date(start, "start"), "to": check_date(end, "end"), "order": "start"}
        events = ctx.export(f"categ/{int(category_id)}", params)
    except (ValueError, RuntimeError) as exc:
        return str(exc)
    needle = fold(query.strip())
    note = ""
    if needle and events:
        def label(e: dict) -> str:
            return f"{e.get('title', '')} {e.get('category', '')}"

        exact = [e for e in events if needle in fold(label(e))]
        if exact:
            events = exact
        else:
            events = [e for e in events if proxy_score(query, label(e)) >= MIN_PROXY]
            note = f"No meeting is named {query!r}. Closest proxies by title or category:\n"
    if not events:
        return f"no meetings found. {ctx.auth_hint()}" if not ctx.token else "no meetings found"
    lines = [
        f"{e['id']} | {when(e)} | {e.get('category', '?')} | {e.get('title', '?')} | {e.get('url', '')}"
        for e in events[-MAX_EVENTS:]
    ]
    if len(events) > MAX_EVENTS:
        lines.insert(0, f"... {len(events)} meetings, showing the latest {MAX_EVENTS}")
    return note + "\n".join(lines)


def tool_get_meeting(ctx: Context, event: str) -> str:
    log(f"get_meeting {event}")
    try:
        ev = load_event(ctx, event)
    except (ValueError, RuntimeError) as exc:
        return str(exc)
    chairs = ", ".join(person(p) for p in ev.get("chairs") or [])
    lines = [
        f"{ev.get('title', '?')}",
        f"category: {ev.get('category', '?')} ({ev.get('categoryId', '?')})",
        f"date: {when(ev)} to {(ev.get('endDate') or {}).get('date', '?')}",
        f"url: {ev.get('url', '')}",
    ]
    if chairs:
        lines.append(f"chairs: {chairs}")
    if ev.get("description"):
        lines.append(f"description: {clip(ev['description'], 1000)}")
    files = attachments(ev)
    if files:
        lines += ["meeting attachments:"] + files
    contribs = sorted(ev.get("contributions") or [], key=when)
    lines.append(f"contributions ({len(contribs)}):")
    lines += [format_contribution(c) for c in contribs]
    return "\n".join(lines)


def tool_find_contributions(
    ctx: Context, name: str, event: str, category_id: int, start: str, end: str
) -> str:
    log(f"find_contributions {name!r} event={event!r} {category_id} {start}..{end}")
    tokens = fold(name).split()
    if not tokens:
        return "name is empty"
    try:
        if event.strip():
            events = [load_event(ctx, event)]
        else:
            params = {
                "from": check_date(start, "start"),
                "to": check_date(end, "end"),
                "order": "start",
                "detail": "contributions",
            }
            events = ctx.export(f"categ/{int(category_id)}", params)
    except (ValueError, RuntimeError) as exc:
        return str(exc)

    hits: list[str] = []
    total = 0
    for ev in events:
        for contrib in sorted(ev.get("contributions") or [], key=when):
            roles = sorted({
                role for role, p in contribution_people(contrib)
                if all(t in fold(f"{p.get('first_name', '')} {p.get('last_name', '')}") for t in tokens)
            })
            if not roles:
                continue
            total += 1
            if len(hits) < MAX_HITS:
                hits.append(
                    f"meeting: {ev.get('title', '?')} ({when(ev)}) {ev.get('url', '')}\n"
                    f"  as {', '.join(roles)}\n{format_contribution(contrib)}"
                )
    if not hits:
        return f"no contributions by {name!r} in {len(events)} meeting(s)"
    body = "\n".join(hits)
    if total > MAX_HITS:
        body += f"\n... {total} contributions, showing {MAX_HITS}"
    return f"{total} contribution(s) in {len(events)} meeting(s) searched\n{body}"


def extract_pages(data: bytes, name: str) -> list[str]:
    """Text per page (PDF), per slide (PPTX), or one page for plain text."""
    lower = name.lower()
    if data[:4] == b"%PDF" or lower.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [page.extract_text() or "" for page in reader.pages]
    if lower.endswith(".pptx"):
        from pptx import Presentation

        pages = []
        for slide in Presentation(io.BytesIO(data)).slides:
            texts = [
                shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
            ]
            pages.append("\n".join(texts))
        return pages
    if lower.endswith((".txt", ".md", ".csv", ".tex", ".html")):
        return [data.decode("utf-8", errors="replace")]
    raise RuntimeError(f"cannot extract text from {name} (supported: pdf, pptx, txt)")


def attachment_hint(ctx: Context, url: str) -> str:
    """Tell a token without the read:everything scope from a file the user cannot access."""
    if not ctx.token:
        return ctx.auth_hint()
    match = re.match(re.escape(INDICO) + r"/event/\d+/", url)
    if not match:
        return ctx.auth_hint()
    try:
        page = ctx.session.get(match.group(0), timeout=60, allow_redirects=False)
    except requests.RequestException:
        return ctx.auth_hint()
    if page.status_code != 200:
        return (
            "The Indico token cannot open web pages: it needs the 'Everything (only GET)' "
            "(read:everything) scope to download slides. Listing meetings only needs "
            "'Classic API (read only)'. Edit the token under My profile > Personal Tokens."
        )
    return "The token works, but this attachment is restricted and your account cannot access it."


def tool_read_attachment(ctx: Context, url: str, start_page: int, end_page: int) -> str:
    log(f"read_attachment {url} {start_page}-{end_page}")
    url = url.strip()
    if not url.startswith(INDICO + "/"):
        return f"Refusing to download {url!r}: only {INDICO} attachments are allowed"
    if url not in ctx.docs:
        dest = ctx.cache / hashlib.sha1(url.encode()).hexdigest()
        try:
            resp = ctx.session.get(url, timeout=300)
        except requests.RequestException as exc:
            return f"download failed: {exc}"
        if resp.status_code in (401, 403):
            return f"download returned {resp.status_code}. {attachment_hint(ctx, url)}"
        if resp.status_code != 200:
            return f"download returned {resp.status_code}. {ctx.auth_hint()}"
        if "text/html" in resp.headers.get("content-type", "") and not url.endswith(".html"):
            return f"Indico returned a web page instead of the file. {ctx.auth_hint()}"
        if len(resp.content) > MAX_BYTES:
            return f"attachment is too large ({len(resp.content)} bytes)"
        dest.write_bytes(resp.content)
        try:
            ctx.docs[url] = extract_pages(resp.content, url.rsplit("/", 1)[-1])
        except Exception as exc:  # broken PDFs raise a variety of errors
            return f"failed to extract text: {exc}"
    pages = ctx.docs[url]
    n = len(pages)
    if n == 0:
        return "(no pages)"
    start = max(1, start_page)
    end = min(n, end_page if end_page else n, start + MAX_PAGES - 1)
    if start > n:
        return f"start_page {start} is past the end ({n} pages)"
    parts = [f"{n} pages, showing {start}-{end}"]
    used = 0
    for i in range(start, end + 1):
        text = pages[i - 1].strip() or "(no text on this page)"
        if used + len(text) > MAX_CHARS:
            parts.append(f"... stopped at page {i - 1}, character limit reached")
            break
        used += len(text)
        parts.append(f"--- page {i} ---\n{text}")
    return "\n".join(parts)


def load_config(path: Path) -> dict[str, str]:
    try:
        data = json.loads(path.read_text())
    except json.JSONDecodeError as exc:
        raise SystemExit(f"Invalid JSON in {path}: {exc}") from exc
    if not isinstance(data, dict):
        raise SystemExit(f"{path} must contain a JSON object")
    cfg: dict[str, str] = {}
    for key in ("api_key", "model"):
        value = data.get(key)
        if not isinstance(value, str) or not value.strip():
            raise SystemExit(f"{path} must contain a non-empty {key}")
        cfg[key] = value.strip()
    for key in ("base_url", "indico_token"):
        value = data.get(key)
        if value is None:
            continue
        if not isinstance(value, str) or not value.strip():
            raise SystemExit(f"{path} {key} must be a non-empty string")
        cfg[key] = value.strip().rstrip("/") if key == "base_url" else value.strip()
    return cfg


def build_model(cfg: dict[str, str]):
    if cfg.get("base_url"):
        return OpenAIChatModel(
            cfg["model"],
            provider=OpenAIProvider(base_url=cfg["base_url"], api_key=cfg["api_key"]),
            profile=OpenAIModelProfile(openai_supports_strict_tool_definition=False),
        )
    return OpenRouterModel(
        cfg["model"],
        provider=OpenRouterProvider(api_key=cfg["api_key"]),
    )


def build_agent(ctx: Context, cfg: dict[str, str], category: int) -> Agent:
    today = datetime.date.today().isoformat()
    context = f" Today is {today}. Default category: {category}."
    agent = Agent(build_model(cfg), instructions=INSTRUCTIONS + context)

    @agent.tool_plain
    def find_category(query: str) -> str:
        """Search Indico categories by name (e.g. 'PWG-HF', 'Physics Forum'). Returns ALICE categories with ids."""
        return tool_find_category(ctx, query)

    @agent.tool_plain
    def list_meetings(category_id: int, start: str, end: str, query: str = "") -> str:
        """List meetings in a category and its subcategories between start and end
        (today, -30d, +14d or YYYY-MM-DD). query filters by meeting title or category name."""
        return tool_list_meetings(ctx, category_id, start, end, query)

    @agent.tool_plain
    def get_meeting(event: str) -> str:
        """Timetable of one meeting (event id or URL): contributions, speakers, and attachment URLs."""
        return tool_get_meeting(ctx, event)

    @agent.tool_plain
    def find_contributions(
        name: str, event: str = "", category_id: int = category, start: str = "-90d", end: str = "+30d"
    ) -> str:
        """Find contributions where a person is speaker or author. If event (id or URL) is given,
        search only that meeting; otherwise search the category between start and end."""
        return tool_find_contributions(ctx, name, event, category_id, start, end)

    @agent.tool_plain
    def read_attachment(url: str, start_page: int = 1, end_page: int = 0) -> str:
        """Extract text from an Indico attachment (PDF or PPTX slides), up to 15 pages per call.
        end_page 0 means as far as the limit allows."""
        return tool_read_attachment(ctx, url, start_page, end_page)

    return agent


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Answer questions about ALICE meetings on Indico with an AI agent.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            '  python alicemeetingassistant.py "What\'s going on in the ALICE Physics Forum?"\n'
            '  python alicemeetingassistant.py "Contributions of Jane Doe in '
            'https://indico.cern.ch/event/1234567/"\n'
            "  python alicemeetingassistant.py            # interactive chat"
        ),
    )
    parser.add_argument(
        "question", nargs="?", help="Question about ALICE meetings (omit to open the chat)"
    )
    parser.add_argument(
        "--config",
        default=None,
        metavar="FILE",
        help="JSON file with api_key, model, optional base_url and indico_token "
        "(default: alicemeetingassistant.json next to this script)",
    )
    parser.add_argument(
        "--category",
        type=int,
        default=ALICE_MEETINGS,
        metavar="ID",
        help=f"Default Indico category to search (default: {ALICE_MEETINGS}, ALICE Meetings)",
    )
    parser.add_argument(
        "--chat", action="store_true", help="After answering the question, continue in the chat"
    )
    args = parser.parse_args()

    default_cfg = Path(SCRIPT_DIR) / "alicemeetingassistant.json"
    cfg_path = Path(args.config).expanduser() if args.config else default_cfg
    if not cfg_path.is_file():
        raise SystemExit(
            f"Missing {cfg_path}. Copy alicemeetingassistant.json.example and set api_key and model."
        )
    cfg = load_config(cfg_path)
    token = os.environ.get("INDICO_TOKEN") or cfg.get("indico_token")
    if not token:
        print("Warning: no indico_token or INDICO_TOKEN set; only public meetings are visible.")
    ctx = Context(token)
    agent = build_agent(ctx, cfg, args.category)
    via = cfg.get("base_url", "OpenRouter")
    print(f"Using {cfg['model']} via {via}")
    limits = UsageLimits(request_limit=50)

    history = None
    if args.question:
        try:
            result = agent.run_sync(args.question, usage_limits=limits)
        except ModelHTTPError as exc:
            detail = exc.body.get("message") if isinstance(exc.body, dict) else None
            detail = detail or exc.message
            hint = f" Try {exc.suggested_model_id!r}." if exc.suggested_model_id else ""
            raise SystemExit(f"{via} rejected {cfg['model']}: {detail}.{hint}") from None
        except UsageLimitExceeded as exc:
            raise SystemExit(str(exc)) from None
        print(result.output)
        if not args.chat:
            return
        history = result.all_messages()

    # The chat UI shows tool calls itself; plain progress lines would break its live view.
    global QUIET
    QUIET = True
    print(
        'Ask about ALICE meetings, e.g. "What\'s going on in the Physics Forum?".\n'
        "Commands: /multiline, /cp (copy last answer), /usage, /exit\n"
    )
    agent.to_cli_sync(prog_name="alice-meetings", message_history=history, usage_limits=limits)


if __name__ == "__main__":
    main()
